# -*- coding: utf-8 -*-
"""Regenere une copie du classeur a partir de docs/data.json.

Le classeur n'est plus le moteur : c'est une photographie, en valeurs, de ce que
le site affiche. Il sert a regarder les chiffres dans Excel, pas a les calculer.
"""
import json
import os
import sys

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SORTIE = os.path.join(RACINE, "classeur", "Analyse_football_paris_sportifs.xlsx")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl absent : copie du classeur ignoree.")
    sys.exit(0)


def feuille(wb, titre, entetes, lignes, notes=()):
    ws = wb.create_sheet(titre)
    ligne = 1
    for n in notes:
        ws.cell(ligne, 1, n).font = Font(italic=True, size=9)
        ligne += 1
    if notes:
        ligne += 1
    for c, h in enumerate(entetes, 1):
        cell = ws.cell(ligne, c, h)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="left")
    for r, row in enumerate(lignes, ligne + 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    for c, h in enumerate(entetes, 1):
        largeur = max([len(str(h))] + [len(str(l[c - 1])) for l in lignes[:200] if l[c - 1] is not None] or [8])
        ws.column_dimensions[get_column_letter(c)].width = min(max(largeur + 2, 9), 34)
    ws.freeze_panes = ws.cell(ligne + 1, 1)
    return ws


def main():
    with open(os.path.join(RACINE, "docs", "data.json"), encoding="utf-8") as f:
        d = json.load(f)
    meta = d.get("meta", {})
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Notice")
    for i, t in enumerate([
        "Analyse football et detection de value bets",
        "",
        "Copie automatique du %s. Ce fichier ne calcule rien : il reprend en valeurs" % meta.get("maj", ""),
        "ce que le site affiche. Le moteur est le depot GitHub, pas ce classeur.",
        "",
        "Fenetre : %s au %s. Competitions : %s." % (
            meta.get("debut", ""), meta.get("fin", ""), ", ".join(meta.get("competitions", []))),
    ], 2):
        ws.cell(i, 2, t)
    ws.column_dimensions["B"].width = 110

    feuille(wb, "Parametres", ["Parametre", "Valeur"],
            [[k, v] for k, v in (meta.get("params") or {}).items()],
            ["Reglages utilises lors du dernier calcul. Les modifier ici ne change rien :",
             "ils vivent dans data/params.json du depot."])

    feuille(wb, "Matchs semaine",
            ["Date", "Heure", "Competition", "Domicile", "Exterieur", "Buts att. dom.",
             "Buts att. ext.", "P(1)", "P(N)", "P(2)", "Verdict", "Confiance", "Score probable",
             "P(+1,5)", "P(+2,5)", "P(BTTS)", "Cote 1", "Cote N", "Cote 2", "Pari", "Edge",
             "Divergence", "Verdict pari", "Source des cotes"],
            [[u.get("date"), u.get("heure"), u.get("comp"), u.get("dom"), u.get("ext"),
              u.get("xd"), u.get("xe"), u.get("p1"), u.get("pn"), u.get("p2"), u.get("verdict"),
              u.get("conf"), u.get("score"), u.get("o15"), u.get("o25"), u.get("btts"),
              u.get("c1"), u.get("cN"), u.get("c2"), u.get("pari"), u.get("edge"),
              u.get("div"), u.get("vp"), u.get("src")] for u in d.get("upcoming", [])])

    feuille(wb, "Selections",
            ["Date", "Heure", "Competition", "Rencontre", "Selection", "Probabilite",
             "Cote estimee", "Cote equitable", "Edge", "Mise conseillee", "Verdict"],
            [[s.get("date"), s.get("heure"), s.get("comp"), s.get("match"), s.get("lab"),
              s.get("p"), s.get("cote"), s.get("fair"), s.get("edge"), s.get("mise"),
              s.get("verdict")] for s in d.get("sel", [])])

    feuille(wb, "Resultats",
            ["Date", "Competition", "Domicile", "Exterieur", "Buts dom.", "Buts ext.", "Resultat",
             "Buts att. dom.", "Buts att. ext.", "P(1)", "P(N)", "P(2)", "Verdict modele",
             "Correct", "Verdict marche", "Correct marche", "Pari", "Cote", "Verdict pari",
             "Resultat pari", "Profit"],
            [[r.get("date"), r.get("comp"), r.get("dom"), r.get("ext"), r.get("bd"), r.get("be"),
              r.get("res"), r.get("xd"), r.get("xe"), r.get("p1"), r.get("pn"), r.get("p2"),
              r.get("vm"), r.get("ok"), r.get("vmar"), r.get("okm"), r.get("pari"), r.get("cote"),
              r.get("vp"), r.get("rp"), r.get("profit")] for r in d.get("results", [])])

    feuille(wb, "Bilan", ["Indicateur", "Valeur"],
            [[k, v] for k, v in (d.get("syn") or {}).items()])
    feuille(wb, "Calibration", ["Tranche", "Matchs", "Reussite observee", "Proba annoncee", "Ecart"],
            d.get("cal", []))
    feuille(wb, "Rendement", ["Verdict", "Paris", "Gagnes", "Reussite", "Profit (1 u)", "ROI",
                              "Esperance marche"], d.get("rend", []))
    feuille(wb, "Par championnat", ["Championnat", "Matchs", "Verdict correct", "Log-loss modele",
                                    "Log-loss marche", "Buts observes", "Buts attendus"],
            d.get("champ", []))

    os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
    wb.save(SORTIE)
    print("classeur regenere :", SORTIE)


if __name__ == "__main__":
    main()
